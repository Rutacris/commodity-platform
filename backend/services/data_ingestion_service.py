import csv
import io
import re
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from typing import Any, Dict, List, Optional, Tuple

import requests
from openpyxl import load_workbook
from sqlalchemy.orm import Session

from models.price import Price


class DataIngestionService:
    @staticmethod
    def parse_records_from_text(text: str, source_label: Optional[str] = None) -> List[Dict[str, str]]:
        cleaned = (text or "").strip()
        if not cleaned:
            return []

        lines = [line.strip() for line in cleaned.splitlines() if line.strip()]
        if not lines:
            return []

        sample = "\n".join(lines[:10])
        delimiter = DataIngestionService._detect_delimiter(sample)
        rows: List[List[str]] = []
        if delimiter:
            rows = list(csv.reader(io.StringIO(sample), delimiter=delimiter))
        else:
            rows = [re.split(r"\s{2,}|\t", line) for line in lines]

        if not rows:
            return []

        header = [cell.strip().lower() for cell in rows[0]]
        data_rows = rows[1:] if DataIngestionService._looks_like_header(header) else rows

        if DataIngestionService._looks_like_header(header):
            field_map = DataIngestionService._build_field_map(header)
            records = []
            for row in data_rows:
                if not row or all(not str(cell).strip() for cell in row):
                    continue
                record = {
                    "commodity": DataIngestionService._extract_value(row, field_map.get("commodity")),
                    "market": DataIngestionService._extract_value(row, field_map.get("market")),
                    "price": DataIngestionService._extract_value(row, field_map.get("price")),
                    "date": DataIngestionService._extract_value(row, field_map.get("date")),
                }
                if record["commodity"] or record["market"] or record["price"]:
                    record["source"] = source_label or "manual"
                    records.append(record)
            return records

        records = []
        for row in data_rows:
            if not row or all(not str(cell).strip() for cell in row):
                continue
            record = {
                "commodity": (row[0].strip() if len(row) > 0 else ""),
                "market": (row[1].strip() if len(row) > 1 else ""),
                "price": (row[2].strip() if len(row) > 2 else ""),
                "date": (row[3].strip() if len(row) > 3 else ""),
            }
            if record["commodity"] or record["market"] or record["price"]:
                record["source"] = source_label or "manual"
                records.append(record)
        return records

    @staticmethod
    def parse_records_from_bytes(filename: str, content_bytes: bytes, source_label: Optional[str] = None) -> List[Dict[str, str]]:
        name = (filename or "").lower()
        if name.endswith(".csv"):
            return DataIngestionService.parse_records_from_text(content_bytes.decode("utf-8-sig", errors="ignore"), source_label=source_label)
        if name.endswith((".xls", ".xlsx")):
            workbook = load_workbook(io.BytesIO(content_bytes), data_only=True)
            sheet = workbook.active
            rows = list(sheet.iter_rows(values_only=True))
            return DataIngestionService._rows_to_records(rows, source_label=source_label)
        return DataIngestionService.parse_records_from_text(content_bytes.decode("utf-8", errors="ignore"), source_label=source_label)

    @staticmethod
    def parse_records_from_url(url: str, source_label: Optional[str] = None) -> List[Dict[str, str]]:
        response = requests.get(url, timeout=20)
        response.raise_for_status()
        name = DataIngestionService._infer_filename_from_url(url, response)
        return DataIngestionService.parse_records_from_bytes(name, response.content, source_label=source_label or f"url:{url}")

    @staticmethod
    def ingest_records(db: Session, records: List[Dict[str, str]], source_label: Optional[str] = None) -> Dict[str, int]:
        inserted = 0
        skipped = 0
        for raw in records:
            normalized = DataIngestionService._normalize_record(raw, source_label=source_label)
            if not normalized:
                continue
            existing = (
                db.query(Price)
                .filter(
                    Price.commodity == normalized["commodity"],
                    Price.market == normalized["market"],
                    Price.date == normalized["date"],
                    Price.source == normalized["source"],
                )
                .first()
            )
            if existing:
                skipped += 1
                continue
            try:
                price_value = Decimal(str(normalized["price"]))
            except (InvalidOperation, ValueError):
                skipped += 1
                continue
            db.add(
                Price(
                    commodity=normalized["commodity"],
                    market=normalized["market"],
                    price=price_value,
                    currency=normalized.get("currency", "RWF"),
                    unit=normalized.get("unit", "kg"),
                    date=normalized["date"],
                    source=normalized["source"],
                )
            )
            inserted += 1

        db.commit()
        return {"inserted": inserted, "skipped": skipped}

    @staticmethod
    def _looks_like_header(header: List[str]) -> bool:
        lowered = [cell.lower().strip() for cell in header if cell]
        return any(name in lowered for name in ["commodity", "market", "price", "date"])

    @staticmethod
    def _build_field_map(header: List[str]) -> Dict[str, int]:
        lowered = [cell.lower().strip() for cell in header]
        field_map: Dict[str, int] = {}
        for idx, name in enumerate(lowered):
            if name in {"commodity", "item", "product", "product_name", "crop"}:
                field_map["commodity"] = idx
            elif name in {"market", "location", "place", "town"}:
                field_map["market"] = idx
            elif name in {"price", "amount", "value", "cost"}:
                field_map["price"] = idx
            elif name in {"date", "record_date", "price_date", "pricedate"}:
                field_map["date"] = idx
        return field_map

    @staticmethod
    def _extract_value(row: List[str], idx: Optional[int]) -> str:
        if idx is None or idx >= len(row):
            return ""
        return str(row[idx]).strip()

    @staticmethod
    def _rows_to_records(rows: List[Tuple[Any, ...]], source_label: Optional[str] = None) -> List[Dict[str, str]]:
        if not rows:
            return []

        header = [str(cell).strip().lower() if cell is not None else "" for cell in rows[0]]
        data_rows = rows[1:] if DataIngestionService._looks_like_header(header) else rows

        if DataIngestionService._looks_like_header(header):
            field_map = DataIngestionService._build_field_map(header)
            records = []
            for row in data_rows:
                record = {
                    "commodity": DataIngestionService._extract_excel_value(row, field_map.get("commodity")),
                    "market": DataIngestionService._extract_excel_value(row, field_map.get("market")),
                    "price": DataIngestionService._extract_excel_value(row, field_map.get("price")),
                    "date": DataIngestionService._extract_excel_value(row, field_map.get("date")),
                }
                if record["commodity"] or record["market"] or record["price"]:
                    record["source"] = source_label or "uploaded"
                    records.append(record)
            return records

        records = []
        for row in data_rows:
            record = {
                "commodity": DataIngestionService._extract_excel_value(row, 0),
                "market": DataIngestionService._extract_excel_value(row, 1),
                "price": DataIngestionService._extract_excel_value(row, 2),
                "date": DataIngestionService._extract_excel_value(row, 3),
            }
            if record["commodity"] or record["market"] or record["price"]:
                record["source"] = source_label or "uploaded"
                records.append(record)
        return records

    @staticmethod
    def _extract_excel_value(row: Tuple[Any, ...], idx: Optional[int]) -> str:
        if idx is None or idx >= len(row):
            return ""
        value = row[idx]
        if value is None:
            return ""
        return str(value).strip()

    @staticmethod
    def _normalize_record(record: Dict[str, str], source_label: Optional[str] = None) -> Optional[Dict[str, Any]]:
        commodity = (record.get("commodity") or "").strip()
        market = (record.get("market") or "").strip()
        price = (record.get("price") or "").strip()
        date_value = (record.get("date") or "").strip()
        if not commodity or not market or not price:
            return None
        if not date_value:
            date_value = date.today().isoformat()
        normalized_date = DataIngestionService._parse_date(date_value)
        if not normalized_date:
            return None
        return {
            "commodity": commodity,
            "market": market,
            "price": price,
            "date": normalized_date,
            "currency": "RWF",
            "unit": "kg",
            "source": (record.get("source") or source_label or "manual"),
        }

    @staticmethod
    def _parse_date(value: str) -> Optional[date]:
        for fmt in ["%Y-%m-%d", "%d/%m/%Y", "%m/%d/%Y", "%Y/%m/%d"]:
            try:
                return datetime.strptime(value, fmt).date()
            except ValueError:
                continue
        return None

    @staticmethod
    def _detect_delimiter(sample: str) -> Optional[str]:
        counts = {",": sample.count(","), "|": sample.count("|"), ";": sample.count(";"), "\t": sample.count("\t")}
        best = max(counts.items(), key=lambda item: item[1])
        if best[1] > 0:
            return best[0]
        return None

    @staticmethod
    def _infer_filename_from_url(url: str, response: requests.Response) -> str:
        disposition = response.headers.get("content-disposition", "")
        if "filename=" in disposition:
            return disposition.split("filename=", 1)[1].strip('"')
        if url.lower().endswith((".csv", ".xls", ".xlsx")):
            return url.split("/")[-1]
        content_type = response.headers.get("content-type", "")
        if "csv" in content_type:
            return "data.csv"
        if "sheet" in content_type or "excel" in content_type:
            return "data.xlsx"
        return "data.txt"
